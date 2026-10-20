"""Text extraction from various document formats."""

import subprocess
from pathlib import Path

from pypdf import PdfReader


def extract_pdf(path: Path) -> str:
    try:
        reader = PdfReader(str(path))
        pages = []
        for page in reader.pages:
            try:
                pages.append(page.extract_text() or "")
            except Exception:
                continue
        return "\n".join(pages)
    except Exception:
        return ""


def extract_docx(path: Path) -> str:
    try:
        import docx

        doc = docx.Document(str(path))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts)
    except Exception:
        return ""


def extract_plain(path: Path) -> str:
    try:
        return path.read_text(errors="ignore")
    except OSError:
        return ""


def extract_rtf_or_doc(path: Path) -> str:
    """Use macOS textutil for .rtf / legacy .doc files."""
    try:
        out = subprocess.run(
            ["textutil", "-convert", "txt", "-stdout", str(path)],
            capture_output=True, text=True, timeout=30,
        )
        if out.returncode == 0:
            return out.stdout
    except Exception:
        pass
    return ""


def extract_text(path: Path) -> str:
    """Route a file to the right extractor; returns '' on failure."""
    suf = path.suffix.lower()
    if suf == ".pdf":
        text = extract_pdf(path)
    elif suf == ".docx":
        text = extract_docx(path)
    elif suf in {".rtf", ".doc"}:
        text = extract_rtf_or_doc(path)
    else:
        text = extract_plain(path)
    return text.strip()


def chunk_text(text: str, chunk_chars: int = 1000, overlap: int = 200) -> list[str]:
    """Split text into overlapping chunks."""
    text = text.strip()
    if not text:
        return []
    if len(text) <= chunk_chars:
        return [text]
    chunks = []
    start = 0
    while start < len(text):
        end = min(start + chunk_chars, len(text))
        chunks.append(text[start:end])
        if end >= len(text):
            break
        start = end - overlap
    return chunks
